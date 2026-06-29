import os
import sys
import time
import logging
import threading
import requests
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font
from flask import Flask, request, jsonify, render_template, redirect, url_for, send_from_directory, session

# Ensure the app_tag_auditor directory is in sys.path
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
if parent_dir not in sys.path:
    sys.path.insert(0, parent_dir)

from core.config import get_settings
from core.drive_client import DriveClient
from core.sheets_client import SheetsClient
from google.oauth2.credentials import Credentials
from orchestrator import run_pipeline

app = Flask(__name__, template_folder='templates')
app.secret_key = os.urandom(24)

# Global variables for tracking the running pipeline
pipeline_thread = None
active_bridge = None
pipeline_error = None
pipeline_status = "idle"  # "idle", "running", "completed", "error"
selected_apk_path = None
selected_schema_path = None
apk_source_info = None
schema_source_info = None

class InteractionBridge:
    def __init__(self):
        self.login_detected = threading.Event()
        self.user_responded = threading.Event()
        self.credentials_needed = threading.Event()
        self.credentials_submitted = threading.Event()
        self.screen_hierarchy = ""
        self.user_decision = None
        self.discovered_fields = []
        self.discovered_escape_options = []
        self.credentials = {}
        self.login_in_progress = False
        self.mid_login_fields = False

class PipelineThread(threading.Thread):
    def __init__(self, apk_path, schema_path, interaction_bridge=None):
        super().__init__()
        self.apk_path = apk_path
        self.schema_path = schema_path
        self.interaction_bridge = interaction_bridge
        self.exception = None

    def run(self):
        global pipeline_status, pipeline_error
        try:
            run_pipeline(
                apk_path=self.apk_path,
                schema_path=self.schema_path,
                interaction_bridge=self.interaction_bridge
            )
            pipeline_status = "completed"
        except Exception as e:
            self.exception = e
            pipeline_error = str(e)
            pipeline_status = "error"

def check_token_refresh():
    """Refreshes the OAuth access token if it is close to expiry."""
    settings = get_settings()
    token_expiry = session.get("google_token_expiry")
    refresh_token = session.get("google_refresh_token")
    if refresh_token and token_expiry:
        if float(token_expiry) < time.time() + 300:
            try:
                res = requests.post(
                    "https://oauth2.googleapis.com/token",
                    data={
                        "client_id": settings.GOOGLE_OAUTH_CLIENT_ID,
                        "client_secret": settings.GOOGLE_OAUTH_CLIENT_SECRET,
                        "refresh_token": refresh_token,
                        "grant_type": "refresh_token",
                    },
                    timeout=10,
                )
                if res.status_code == 200:
                    data = res.json()
                    session["google_auth_token"] = data.get("access_token")
                    session["google_token_expiry"] = str(time.time() + data.get("expires_in", 3600))
            except Exception:
                pass

@app.route("/")
def index():
    settings = get_settings()
    code = request.args.get("code")
    
    # Check if Google returned oauth code in the URL query params
    if code:
        try:
            token_res = requests.post(
                "https://oauth2.googleapis.com/token",
                data={
                    "client_id": settings.GOOGLE_OAUTH_CLIENT_ID,
                    "client_secret": settings.GOOGLE_OAUTH_CLIENT_SECRET,
                    "code": code,
                    "grant_type": "authorization_code",
                    "redirect_uri": settings.GOOGLE_OAUTH_REDIRECT_URI
                },
                timeout=10
            )
            if token_res.status_code == 200:
                token_data = token_res.json()
                session["google_auth_token"] = token_data.get("access_token")
                session["google_refresh_token"] = token_data.get("refresh_token")
                session["google_token_expiry"] = str(time.time() + token_data.get("expires_in", 3600))
                
                profile_res = requests.get(
                    "https://www.googleapis.com/oauth2/v3/userinfo",
                    headers={"Authorization": f"Bearer {session['google_auth_token']}"},
                    timeout=10
                )
                if profile_res.status_code == 200:
                    session["user_profile"] = profile_res.json()
                    session["authenticated"] = True
        except Exception as e:
            logging.error(f"OAuth exchange error: {e}")
        return redirect(url_for("index"))

    return render_template("index.html")

@app.route("/api/profile", methods=["GET"])
def get_profile():
    if not session.get("authenticated"):
        return jsonify({"authenticated": False, "profile": None})
    check_token_refresh()
    return jsonify({
        "authenticated": True,
        "profile": session.get("user_profile"),
        "token": session.get("google_auth_token")
    })

@app.route("/api/config", methods=["GET"])
def get_config():
    settings = get_settings()
    return jsonify({
        "google_client_id": settings.GOOGLE_OAUTH_CLIENT_ID,
        "google_redirect_uri": settings.GOOGLE_OAUTH_REDIRECT_URI
    })

@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"success": True})

@app.route("/api/upload/apk", methods=["POST"])
def upload_apk():
    global selected_apk_path, apk_source_info
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "No file selected"}), 400
    
    settings = get_settings()
    temp_dir = settings.TEMP_STORAGE_DIR
    os.makedirs(temp_dir, exist_ok=True)
    
    dest_path = os.path.join(temp_dir, file.filename)
    file.save(dest_path)
    selected_apk_path = dest_path
    apk_source_info = {"source": "local", "filename": file.filename}
    return jsonify({"success": True, "filename": file.filename})

@app.route("/api/upload/schema", methods=["POST"])
def upload_schema():
    global selected_schema_path, schema_source_info
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "No file selected"}), 400
    
    settings = get_settings()
    temp_dir = settings.TEMP_STORAGE_DIR
    os.makedirs(temp_dir, exist_ok=True)
    
    dest_path = os.path.join(temp_dir, file.filename)
    file.save(dest_path)
    selected_schema_path = dest_path
    schema_source_info = {"source": "local", "filename": file.filename}
    return jsonify({"success": True, "filename": file.filename})

@app.route("/api/drive/picker", methods=["POST"])
def drive_picker():
    global selected_apk_path, apk_source_info
    data = request.json or {}
    file_id = data.get("file_id")
    file_name = data.get("file_name", "drive_file.apk")
    token = session.get("google_auth_token") or data.get("access_token")
    
    if not file_id or not token:
        return jsonify({"error": "Missing file_id or auth token"}), 400
        
    settings = get_settings()
    temp_dir = settings.TEMP_STORAGE_DIR
    os.makedirs(temp_dir, exist_ok=True)
    dest_path = os.path.join(temp_dir, file_name)
    
    try:
        creds = Credentials(token=token)
        drive_client = DriveClient(credentials=creds)
        drive_client.download_file(file_id, dest_path)
        selected_apk_path = dest_path
        apk_source_info = {"source": "drive", "filename": file_name, "file_id": file_id}
        return jsonify({"success": True, "filename": file_name})
    except Exception as e:
        return jsonify({"error": f"Failed to download from Drive: {str(e)}"}), 500

@app.route("/api/sheets/picker", methods=["POST"])
def sheets_picker():
    global selected_schema_path, schema_source_info
    data = request.json or {}
    sheet_url = data.get("sheet_url")
    token = session.get("google_auth_token")
    
    if not sheet_url or not token:
        return jsonify({"error": "Missing sheet_url or authentication"}), 400
        
    settings = get_settings()
    temp_dir = settings.TEMP_STORAGE_DIR
    os.makedirs(temp_dir, exist_ok=True)
    
    try:
        sheet_id = SheetsClient.extract_sheet_id_from_url(sheet_url)
        creds = Credentials(token=token)
        sheets_client = SheetsClient(credentials=creds)
        dest_path = os.path.join(temp_dir, f"sheet_{sheet_id}.xlsx")
        sheets_client.download_sheet_as_excel(sheet_id, dest_path)
        selected_schema_path = dest_path
        schema_source_info = {"source": "sheets", "filename": f"sheet_{sheet_id}.xlsx", "sheet_url": sheet_url}
        return jsonify({"success": True, "filename": f"sheet_{sheet_id}.xlsx"})
    except Exception as e:
        return jsonify({"error": f"Failed to download Google Sheet: {str(e)}"}), 500

@app.route("/api/run", methods=["POST"])
def run():
    global pipeline_thread, active_bridge, pipeline_status, pipeline_error
    if not selected_apk_path or not selected_schema_path:
        return jsonify({"error": "Missing selected APK or schema sheet"}), 400
        
    if pipeline_status == "running":
        return jsonify({"error": "Pipeline is already running"}), 400
        
    # Clear the log file before starting
    log_file = "logs/app_tag_auditor.log"
    if os.path.exists(log_file):
        try:
            with open(log_file, "w") as f:
                f.write("")
        except Exception:
            pass

    active_bridge = InteractionBridge()
    pipeline_error = None
    pipeline_status = "running"
    
    pipeline_thread = PipelineThread(
        apk_path=selected_apk_path,
        schema_path=selected_schema_path,
        interaction_bridge=active_bridge
    )
    pipeline_thread.start()
    
    return jsonify({"success": True})

@app.route("/api/status", methods=["GET"])
def status():
    global pipeline_status, pipeline_error, active_bridge
    
    # Read log lines
    logs = ""
    log_file = "logs/app_tag_auditor.log"
    if os.path.exists(log_file):
        try:
            with open(log_file, "r") as f:
                logs = f.read()
        except Exception:
            pass
            
    hitl = None
    if active_bridge and active_bridge.login_detected.is_set():
        hitl = {
            "mid_login_fields": active_bridge.mid_login_fields,
            "fields": active_bridge.discovered_fields,
            "escape_options": active_bridge.discovered_escape_options
        }
        
    return jsonify({
        "status": pipeline_status,
        "error": pipeline_error,
        "logs": logs,
        "hitl": hitl,
        "apk_info": apk_source_info,
        "schema_info": schema_source_info
    })

@app.route("/api/hitl/respond", methods=["POST"])
def hitl_respond():
    global active_bridge
    if not active_bridge or not active_bridge.login_detected.is_set():
        return jsonify({"error": "No pending intervention prompt"}), 400
        
    data = request.json or {}
    decision = data.get("decision")
    credentials = data.get("credentials", {})
    
    if not decision:
        return jsonify({"error": "Missing decision"}), 400
        
    active_bridge.user_decision = decision
    active_bridge.credentials = credentials
    
    if decision in ["login", "signup", "submit_fields"]:
        active_bridge.login_in_progress = True
    else:
        active_bridge.login_in_progress = False
        
    active_bridge.user_responded.set()
    active_bridge.login_detected.clear()
    
    return jsonify({"success": True})

@app.route("/api/results", methods=["GET"])
def get_results():
    settings = get_settings()
    output_path = settings.LOCAL_OUTPUT_PATH
    if not os.path.exists(output_path):
        return jsonify({"error": "Report not found"}), 404
        
    try:
        xl = pd.ExcelFile(output_path)
        sheet_names = xl.sheet_names
        preferred_order = ["Audit Summary", "Audit Analysis"]
        sorted_sheet_names = [s for s in preferred_order if s in sheet_names] + [s for s in sheet_names if s not in preferred_order]
        
        sheets_data = {}
        for name in sorted_sheet_names:
            if name == "Audit Summary":
                try:
                    df = pd.read_excel(output_path, sheet_name=name, header=3)
                except Exception:
                    df = pd.read_excel(output_path, sheet_name=name)
            else:
                df = pd.read_excel(output_path, sheet_name=name)
            df = df.fillna("")
            sheets_data[name] = {
                "headers": list(df.columns),
                "rows": df.to_dict(orient="records")
            }
        return jsonify({"sheets": sorted_sheet_names, "data": sheets_data})
    except Exception as e:
        return jsonify({"error": f"Failed to read report: {str(e)}"}), 500

@app.route("/api/results/save", methods=["POST"])
def save_results():
    settings = get_settings()
    output_path = settings.LOCAL_OUTPUT_PATH
    if not os.path.exists(output_path):
        return jsonify({"error": "Report not found"}), 404
        
    data = request.json or {}
    active_sheet = data.get("sheet")
    rows = data.get("rows", {})
    
    try:
        wb = openpyxl.load_workbook(output_path)
        if active_sheet in wb.sheetnames:
            ws = wb[active_sheet]
            headers = [cell.value for cell in ws[1]]
            
            try:
                status_col_idx = headers.index("Status") + 1
            except ValueError:
                status_col_idx = None
            try:
                comments_col_idx = headers.index("Comments") + 1
            except ValueError:
                comments_col_idx = None
            try:
                logs_col_idx = headers.index("Logs") + 1
            except ValueError:
                logs_col_idx = None
                
            font_family = "Segoe UI"
            green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            green_font = Font(name=font_family, size=10, bold=True, color="375623")
            yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            yellow_font = Font(name=font_family, size=10, bold=True, color="7F6000")
            red_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
            red_font = Font(name=font_family, size=10, bold=True, color="78281F")
            gray_fill = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")
            gray_font = Font(name=font_family, size=10, bold=True, color="5D6D7E")
            
            for row_str, changes in rows.items():
                row_idx = int(row_str)
                excel_row = row_idx + 2
                
                if "status" in changes and status_col_idx:
                    val = changes["status"]
                    cell = ws.cell(row=excel_row, column=status_col_idx, value=val)
                    if val == "Implemented":
                        cell.fill = green_fill
                        cell.font = green_font
                    elif val == "Implemented with issues":
                        cell.fill = yellow_fill
                        cell.font = yellow_font
                    elif val == "Scenario Not Found":
                        cell.fill = gray_fill
                        cell.font = gray_font
                    else:
                        cell.fill = red_fill
                        cell.font = red_font
                        
                if "comments" in changes and comments_col_idx:
                    ws.cell(row=excel_row, column=comments_col_idx, value=changes["comments"])
                    
                if "logs" in changes and logs_col_idx:
                    ws.cell(row=excel_row, column=logs_col_idx, value=changes["logs"])
            
            # Recalculate summary tab counts
            if "Audit Summary" in wb.sheetnames and status_col_idx:
                ws_summary = wb["Audit Summary"]
                implemented_count = 0
                implemented_with_issues_count = 0
                not_implemented_count = 0
                scenario_not_found_count = 0
                
                for r in range(2, ws.max_row + 1):
                    val = ws.cell(row=r, column=status_col_idx).value
                    if val == "Implemented":
                        implemented_count += 1
                    elif val == "Implemented with issues":
                        implemented_with_issues_count += 1
                    elif val == "Scenario Not Found":
                        scenario_not_found_count += 1
                    elif val is not None:
                        not_implemented_count += 1
                        
                ws_summary["B5"] = implemented_count
                ws_summary["B6"] = implemented_with_issues_count
                ws_summary["B7"] = not_implemented_count
                ws_summary["B8"] = scenario_not_found_count
                
            wb.save(output_path)
            return jsonify({"success": True})
        return jsonify({"error": f"Sheet {active_sheet} not found"}), 404
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/results/download", methods=["GET"])
def download_results():
    settings = get_settings()
    output_path = settings.LOCAL_OUTPUT_PATH
    if not os.path.exists(output_path):
        return jsonify({"error": "Report not found"}), 404
    
    directory = os.path.dirname(os.path.abspath(output_path))
    filename = os.path.basename(output_path)
    
    # Generate download filename dynamically based on input schema name
    given_sheet_name = "results"
    if schema_source_info:
        if schema_source_info["source"] == "sheets" and "sheet_url" in schema_source_info:
            given_sheet_name = SheetsClient.extract_sheet_id_from_url(schema_source_info["sheet_url"])
        else:
            given_sheet_name = os.path.splitext(schema_source_info["filename"])[0]
    
    export_filename = f"audit_result_{given_sheet_name}.xlsx"
    return send_from_directory(directory, filename, as_attachment=True, download_name=export_filename)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8501, debug=True)
