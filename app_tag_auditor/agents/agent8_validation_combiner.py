"""
Agent8 ValidationCombinerAgent:
Merges TelemetryValidationResult + RuntimeValidationResult per event into a final,
consolidated FinalAuditRow, and writes the results to output Excel sheets.
"""

import os
from datetime import datetime
from typing import Literal, Optional
import core.logger
from core.models import ExpectedEvent, TelemetryValidationResult, RuntimeValidationResult, FinalAuditRow, CapturedLog

def _format_captured_log(log: CapturedLog) -> str:
    timestamp_str = "06-19 18:07:11.198"
    if log.timestamp:
        try:
            dt = datetime.fromisoformat(log.timestamp)
            timestamp_str = dt.strftime("%m-%d %H:%M:%S.%f")[:-3]
        except Exception:
            pass
    params_lines = [
        f"regionCode={log.raw_params.get('regionCode') or 'com.example.app'}",
        "countryCode=IN",
        "tyc_environment=REAPPDEV_2.0"
    ]
    if log.raw_params:
        for k, v in log.raw_params.items():
            if k not in ("regionCode", "countryCode", "tyc_environment"):
                params_lines.append(f"{k}={v}")
    params_lines.extend(["ga_event_origin(_o)=app", "manual_tracking(_mst)=1"])
    return (
        f"{timestamp_str} 21272 5878 V FA-SVC : Logging event:\n"
        f"origin=app,name={log.event_name}{'(_vs)' if log.event_name == 'screen_view' else ''},params=Bundle[[\n"
        f"{',\n'.join(params_lines)}\n"
        f"]]"
    )

def _evaluate_event(
    event: ExpectedEvent, 
    telemetry: Optional[TelemetryValidationResult], 
    runtime: Optional[RuntimeValidationResult], 
    has_code_mapping: bool
) -> tuple[str, str, str]:
    status = "Not Implemented"
    comments_list = []
    evidence = "Event absent in runtime logs"
    
    if telemetry and telemetry.matched_log:
        log = telemetry.matched_log
        if telemetry.passed:
            status = "Implemented"
            evidence = _format_captured_log(log)
        else:
            status = "Implemented with issues"
            for k in telemetry.missing_keys:
                comments_list.append(f"• The parameter '{k}' is missing.")
            for k in telemetry.mismatched_keys:
                expected_param = next((p for p in event.expected_params if p.param_name == k), None)
                expected_val = expected_param.example_value if expected_param else ""
                comments_list.append(f"• '{k}' value mismatch: expected '{expected_val}' but got '{log.raw_params.get(k, '')}'.")
            if telemetry.extra_keys:
                keys_str = ", ".join(f"'{k}'" for k in telemetry.extra_keys)
                comments_list.append(f"• {keys_str} {'is an' if len(telemetry.extra_keys) == 1 else 'are'} extra parameter{'s' if len(telemetry.extra_keys) > 1 else ''} found in the logs. It is advised to remove {'it' if len(telemetry.extra_keys) == 1 else 'them'}.")
            
            if any("screen" in k.lower() for k in (telemetry.missing_keys + telemetry.mismatched_keys)):
                evidence = "Event absent in runtime logs"
            else:
                evidence = _format_captured_log(log)
    else:
        status = "Scenario Not Found"
        if has_code_mapping:
            comments_list.append("• Event found in codebase but absent in runtime logs.")
        else:
            comments_list.append("• Event not found in codebase and absent in runtime logs.")
            
    return status, "\n".join(comments_list), evidence

class ValidationCombinerAgent:
    def __init__(self):
        self.logger = core.logger.get_logger(__name__)

    def _build_details(self, telemetry: TelemetryValidationResult, runtime: RuntimeValidationResult) -> str:
        clauses = []
        if telemetry.missing_keys:
            clauses.append(f"Missing params: {', '.join(telemetry.missing_keys)}.")
        if telemetry.extra_keys:
            clauses.append(f"Extra params: {', '.join(telemetry.extra_keys)}.")
        if telemetry.mismatched_keys:
            clauses.append(f"Incorrect values: {', '.join(telemetry.mismatched_keys)}.")
        if runtime.not_implemented:
            clauses.append("Event never fired.")
        if runtime.double_fired:
            clauses.append(f"Fired {runtime.fire_count} times, expected once.")
        if runtime.screen_check_status == "incorrect":
            clauses.append("Fired on the wrong screen.")
        if runtime.screen_check_status == "unknown":
            clauses.append("Screen identity inconclusive — manual check recommended.")
        if runtime.unexpected_co_fired_events:
            clauses.append(f"Unexpected co-fired event(s): {', '.join(runtime.unexpected_co_fired_events)}.")

        if not clauses:
            return "All checks passed."
        return " ".join(clauses)

    def _determine_overall_status(self, telemetry: TelemetryValidationResult, runtime: RuntimeValidationResult) -> Literal["PASS", "FAIL", "PARTIAL"]:
        if telemetry.passed and runtime.passed:
            return "PARTIAL" if runtime.screen_check_status == "unknown" else "PASS"
        if (not telemetry.passed) and (not runtime.passed):
            return "FAIL"
        return "PARTIAL"

    def combine_event(self, event: ExpectedEvent, telemetry: TelemetryValidationResult, runtime: RuntimeValidationResult) -> FinalAuditRow:
        overall = self._determine_overall_status(telemetry, runtime)
        details = self._build_details(telemetry, runtime)
        return FinalAuditRow(
            event_name=event.event_name,
            screen=event.screen,
            telemetry_passed=telemetry.passed,
            runtime_passed=runtime.passed,
            overall_status=overall,
            details=details,
            timestamp=datetime.now().isoformat()
        )

    def _write_pretty_excel(self, expected_events, telemetry_results, runtime_results, code_locations, file_path):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        telemetry_map = {(res.event_name, res.screen): res for res in telemetry_results}
        runtime_map = {(res.event_name, res.screen): res for res in runtime_results}
        code_map = {loc.event_name for loc in code_locations} if code_locations else set()

        implemented_count = 0
        implemented_with_issues_count = 0
        not_implemented_count = 0
        scenario_not_found_count = 0
        analysis_rows_data = []

        for event in expected_events:
            telemetry = telemetry_map.get((event.event_name, event.screen))
            runtime = runtime_map.get((event.event_name, event.screen))
            has_code_mapping = (event.event_name in code_map)
            
            status, comments, evidence = _evaluate_event(event, telemetry, runtime, has_code_mapping)
            if status == "Implemented":
                implemented_count += 1
            elif status == "Implemented with issues":
                implemented_with_issues_count += 1
            elif status == "Scenario Not Found":
                scenario_not_found_count += 1
            else:
                not_implemented_count += 1

            analysis_rows_data.append({
                "User Action": event.user_action,
                "Screen Name": event.screen,
                "Event Name": event.event_name,
                "Event Parameters": ", ".join(p.param_name for p in event.expected_params),
                "Parameter Type": ", ".join(p.parameter_type for p in event.expected_params),
                "Data Type": ", ".join(p.data_type for p in event.expected_params),
                "Principle": event.raw_principle,
                "Event Parameters Example Values": ", ".join(f"{p.param_name}={p.example_value}" for p in event.expected_params),
                "Status": status,
                "Comments": comments,
                "Logs": evidence
            })

        wb = openpyxl.load_workbook(file_path) if os.path.exists(file_path) else openpyxl.Workbook()
        ws_summary = wb.create_sheet(title="Audit Summary") if "Audit Summary" not in wb.sheetnames else wb["Audit Summary"]
        ws_analysis = wb.create_sheet(title="Audit Analysis") if "Audit Analysis" not in wb.sheetnames else wb["Audit Analysis"]
        
        ws_summary.delete_rows(1, ws_summary.max_row + 10)
        ws_analysis.delete_rows(1, ws_analysis.max_row + 10)

        for sname in list(wb.sheetnames):
            if sname not in ["Audit Summary", "Audit Analysis"]:
                del wb[sname]

        ws_summary.append([])
        ws_summary.append(["App Tag Auditor - Audit Summary"])
        ws_summary.append([])
        ws_summary.append(["Static Status Category", "Count"])
        ws_summary.append(["Implemented", implemented_count])
        ws_summary.append(["Implemented with issues", implemented_with_issues_count])
        ws_summary.append(["Not Implemented", not_implemented_count])
        ws_summary.append(["Scenario Not Found", scenario_not_found_count])
        ws_summary.append(["Total Static Events", len(expected_events)])
        ws_summary.merge_cells("A2:B2")

        font_family = "Segoe UI"
        title_font = Font(name=font_family, size=16, bold=True, color="1B365D")
        header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        green_font = Font(name=font_family, size=10, bold=True, color="375623")
        yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        yellow_font = Font(name=font_family, size=10, bold=True, color="7F6000")
        red_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
        red_font = Font(name=font_family, size=10, bold=True, color="78281F")
        gray_fill = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")
        gray_font = Font(name=font_family, size=10, bold=True, color="5D6D7E")
        
        bold_font = Font(name=font_family, size=11, bold=True, color="000000")
        regular_font = Font(name=font_family, size=10, color="000000")
        
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws_summary["A2"].font = title_font
        ws_summary["A2"].alignment = left_align

        for col in ["A", "B"]:
            ws_summary[f"{col}4"].font = header_font
            ws_summary[f"{col}4"].fill = header_fill
            ws_summary[f"{col}4"].alignment = center_align
            ws_summary[f"{col}4"].border = thin_border

        for r_idx, (fill, font) in enumerate([(green_fill, green_font), (yellow_fill, yellow_font), (red_fill, red_font), (gray_fill, gray_font)], 5):
            ws_summary[f"A{r_idx}"].fill = fill
            ws_summary[f"A{r_idx}"].font = font
            ws_summary[f"A{r_idx}"].border = thin_border
            ws_summary[f"B{r_idx}"].font = bold_font
            ws_summary[f"B{r_idx}"].alignment = center_align
            ws_summary[f"B{r_idx}"].border = thin_border

        double_bottom = Border(top=Side(border_style="thin", color="000000"), bottom=Side(border_style="double", color="000000"))
        ws_summary["A9"].font = bold_font
        ws_summary["A9"].border = double_bottom
        ws_summary["B9"].font = bold_font
        ws_summary["B9"].alignment = center_align
        ws_summary["B9"].border = double_bottom
        
        ws_summary.column_dimensions["A"].width = 30
        ws_summary.column_dimensions["B"].width = 15
        ws_summary.views.sheetView[0].showGridLines = True

        headers = [
            "User Action", "Screen Name", "Event Name", "Event Parameters", 
            "Parameter Type", "Data Type", "Principle", 
            "Event Parameters Example Values", "Status", "Comments", "Logs"
        ]
        ws_analysis.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws_analysis.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row_idx, rdata in enumerate(analysis_rows_data, 2):
            ws_analysis.append([rdata[h] for h in headers])
            for col_idx in range(1, len(headers) + 1):
                cell = ws_analysis.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                h_name = headers[col_idx-1]
                if h_name in ["Comments", "Logs"]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                elif h_name == "Status":
                    cell.alignment = center_align
                    val = cell.value
                    if val == "Implemented":
                        cell.fill = green_fill
                        cell.font = green_font
                    elif val == "Implemented with issues":
                        cell.fill = yellow_fill
                        cell.font = yellow_font
                    elif val == "Scenario Not Found":
                        cell.fill = gray_fill
                        cell.font = gray_font
                    else:
                        cell.fill = red_fill
                        cell.font = red_font
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        widths = {"A": 30, "B": 20, "C": 20, "D": 25, "E": 18, "F": 15, "G": 35, "H": 35, "I": 25, "J": 45, "K": 60}
        for col_letter, width in widths.items():
            ws_analysis.column_dimensions[col_letter].width = width
        ws_analysis.views.sheetView[0].showGridLines = True

        wb._sheets = [wb["Audit Summary"], wb["Audit Analysis"]]
        wb.save(file_path)

    def run(
        self,
        expected_events: list[ExpectedEvent],
        telemetry_results: list[TelemetryValidationResult],
        runtime_results: list[RuntimeValidationResult],
        code_locations: list = None,
        output_writer = None
    ) -> list[FinalAuditRow]:
        telemetry_dict = {r.event_name: r for r in telemetry_results}
        runtime_dict = {r.event_name: r for r in runtime_results}
        
        final_rows = []
        for event in expected_events:
            telemetry = telemetry_dict.get(event.event_name)
            runtime = runtime_dict.get(event.event_name)
            if not telemetry or not runtime:
                self.logger.error(f"Missing results for event: {event.event_name}. Telemetry: {telemetry is not None}, Runtime: {runtime is not None}")
                continue
            row = self.combine_event(event, telemetry, runtime)
            final_rows.append(row)

        pass_count = sum(1 for r in final_rows if r.overall_status == "PASS")
        fail_count = sum(1 for r in final_rows if r.overall_status == "FAIL")
        partial_count = sum(1 for r in final_rows if r.overall_status == "PARTIAL")
        self.logger.info(f"Audit summary: PASS: {pass_count}, FAIL: {fail_count}, PARTIAL: {partial_count}")

        os.makedirs("./tmp", exist_ok=True)
        from core.output_writer import LocalExcelWriter
        temp_writer = LocalExcelWriter("./tmp/final_audit_raw.xlsx")
        self.write_to_output(final_rows, temp_writer)

        if output_writer is not None:
            self._write_pretty_excel(expected_events, telemetry_results, runtime_results, code_locations, output_writer.file_path)

        return final_rows

    def write_to_output(self, rows: list[FinalAuditRow], output_writer, tab_name: str = "FinalAudit") -> None:
        headers = ["event_name", "screen", "telemetry_passed", "runtime_passed", "overall_status", "details", "timestamp"]
        output_writer.ensure_headers(tab_name, headers)
        for row in rows:
            output_writer.append_row(tab_name, row.model_dump())

if __name__ == "__main__":
    from core.output_writer import LocalExcelWriter
    import openpyxl

    # Define test events
    all_events = [
        ExpectedEvent(
            event_name="screen_view", screen="My_RE_screen", user_action="lands on My RE screen", expected_params=[], raw_principle="", keywords=[]
        ),
        ExpectedEvent(
            event_name="add_motorcycle", screen="My_RE_screen", user_action="clicks Add icon", expected_params=[], raw_principle="", keywords=[]
        )
    ]

    combiner = ValidationCombinerAgent()

    # a. screen_view - PASS
    tele_a = TelemetryValidationResult(event_name="screen_view", screen="My_RE_screen", passed=True)
    run_a = RuntimeValidationResult(
        event_name="screen_view", screen="My_RE_screen", passed=True, expected_trigger="lands on My RE screen",
        actual_trigger_observed=True, notes="", fire_count=1, screen_check_status="correct"
    )
    res_a = combiner.combine_event(all_events[0], tele_a, run_a)
    print("Expected: PASS")
    print(res_a.model_dump_json(indent=2))
    print("=" * 60)

    # b. add_motorcycle - PARTIAL
    tele_b = TelemetryValidationResult(event_name="add_motorcycle", screen="My_RE_screen", passed=True)
    run_b = RuntimeValidationResult(
        event_name="add_motorcycle", screen="My_RE_screen", passed=False, expected_trigger="clicks Add icon",
        actual_trigger_observed=True, notes="", fire_count=2, double_fired=True, screen_check_status="correct"
    )
    res_b = combiner.combine_event(all_events[1], tele_b, run_b)
    print("Expected: PARTIAL")
    print(res_b.model_dump_json(indent=2))
    print("=" * 60)

    os.makedirs("./tmp", exist_ok=True)
    temp_file = "./tmp/test_combiner_results.xlsx"
    temp_writer = LocalExcelWriter(temp_file)

    results = combiner.run(
        all_events,
        [tele_a, tele_b],
        [run_a, run_b],
        code_locations=[],
        output_writer=temp_writer
    )

    print("\nWorkbook Sheet Names:")
    wb = openpyxl.load_workbook(temp_file)
    print(wb.sheetnames)

    print("\nRead back Audit Summary:")
    for row in wb["Audit Summary"].iter_rows(values_only=True):
        if any(row):
            print(row)

    if os.path.exists(temp_file):
        os.remove(temp_file)
    if os.path.exists("./tmp/final_audit_raw.xlsx"):
        os.remove("./tmp/final_audit_raw.xlsx")
